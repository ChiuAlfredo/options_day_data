import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import get_column_letter
import os

# 讀取Excel文件
df = pd.read_excel("作業.xlsx")

# 將 '時間' 列轉換為datetime格式，只提取時間部分
df['Time'] = pd.to_datetime(df['時間'], errors='coerce').dt.time

# 定義時間範圍
start_time = pd.to_datetime('09:00:00').time()
end_time = pd.to_datetime('13:30:00').time()

# 篩選時間範圍
timefiltered_df = df[(df['Time'] >= start_time) & (df['Time'] <= end_time)]

# 設置支持中文的字體
plt.rcParams['font.sans-serif'] = ['SimHei']  # 使用黑體字
plt.rcParams['axes.unicode_minus'] = False  # 解決座標軸負號顯示問題

# 創建散佈圖
plt.figure(figsize=(10, 6))
plt.scatter(timefiltered_df['期貨_成交價 變化比'], timefiltered_df['發行量加權股價指數 變化比'], alpha=0.6)
plt.title('臺指期成交價變化率與加權指數變化率散佈圖', fontsize=25)
plt.xlabel('Delta F/F', fontsize=20)
plt.ylabel('Delta S/S', fontsize=20)
plt.grid(True)

# 指定保存路徑
save_path = "C:/Users/ASUS/Desktop/散佈圖/作業/期貨變化率&加權指數變化率.png"
plt.savefig(save_path)
